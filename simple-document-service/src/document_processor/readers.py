import csv
import logging
import tempfile
from io import BytesIO
from typing import BinaryIO

import chardet
import pymupdf as fitz
import pymupdf4llm
import torch
import xlrd
from PIL import Image
from doclayout_yolo import YOLOv10
from doclayout_yolo.engine.results import Results
from docx import Document
from fastapi import HTTPException
from openpyxl import load_workbook
from optimum.onnxruntime import ORTModelForVision2Seq
from pptx import Presentation
from spire.doc import Document as SpireDocument
from transformers import TrOCRProcessor

from config import settings
from document_processor.misc import estimate_text_density, rectangles_intersect, find_partial_match_end_index

reader = settings.ocr_reader


class PickleableHTTPException(HTTPException):
    def __reduce__(self):
        return self.__class__, (self.status_code, self.detail, self.headers)

layout_detector = YOLOv10(
    r"./models/layout_detector/doclayout_yolo_docstructbench_imgsz1024.pt", verbose=False)

processor = TrOCRProcessor.from_pretrained('./models/formula_recognizer', local_files_only=True, use_fast=True)
formula_recognizer = ORTModelForVision2Seq.from_pretrained('./models/formula_recognizer', use_cache=False, local_files_only=True)


def pdf_reader(pdf_bytes: bytes) -> str:
    try:
        result_text = []
        with fitz.open(stream=pdf_bytes) as doc:
            images_char_count = estimate_text_density(doc)
            if images_char_count >= 80_000: # слишком долго обрабатывать - raise ошибки
                raise PickleableHTTPException(
                    status_code=400,
                    detail="Document is too large",
                    headers={
                        "content_type": "application/pdf",
                        "chars": images_char_count
                    }
                )

            for page in doc:
                page: fitz.Page
                page_text = ''

                page_dict = page.get_text("dict")

                # скриншот страницы
                pix = page.get_pixmap()
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)

                # layout детектор
                det_res: list[Results] = layout_detector.predict(
                    img,  # Image to predict
                    conf=0.2,  # Confidence threshold
                    device="cpu",  # Device to use (e.g., 'cuda:0' or 'cpu'),
                    verbose=False
                )

                # распознавание формул
                classes: torch.Tensor = det_res[0].boxes.cls
                indices_formulas = (classes == 8).nonzero(as_tuple=True)[0]  # фильтрация по классу формул
                formula_list: list[dict] = []
                for formula_index in indices_formulas:
                    formula_coordinates = det_res[0].boxes.xyxy[formula_index]
                    formula_coordinates = formula_coordinates.int().tolist()
                    formula_image = img.crop(formula_coordinates)

                    pixel_values = processor(formula_image, return_tensors="pt").pixel_values
                    generated_ids = formula_recognizer.generate(pixel_values)
                    latex_text = processor.batch_decode(generated_ids, skip_special_tokens=True)[0]

                    formula_list.append({"coordinates": formula_coordinates, "latex": latex_text})

                page_non_interact_text = [[]]  # список для хранения текста, не пересекающегося с формулами
                index = 0

                for block in page_dict["blocks"]:
                    if block["type"] == 0:  # текстовый блок
                        # Проходим по линиям блока
                        for line in block.get("lines", []):
                            # Проходим по каждому span (фрагменту текста)
                            spans: list = line.get("spans", [])
                            for span in spans:
                                bbox = span["bbox"]

                                append_allowed = True
                                for formula in formula_list:
                                    if rectangles_intersect(bbox, formula["coordinates"]):  # если bbox пересекается с формулой
                                        append_allowed = False
                                        index += 1
                                        page_non_interact_text.append([])
                                        break

                                if append_allowed:
                                    page_non_interact_text[index].append(span["text"])

                original_text = pymupdf4llm.to_markdown(doc=doc, pages=[page.number], show_progress=False)

                # удаляем пустые строки и пробелы
                for i in range(len(page_non_interact_text) - 1, -1, -1):
                    sentence: list[str] = page_non_interact_text[i]
                    if not sentence:
                        page_non_interact_text.pop(i)
                    else:
                        for index in range(len(sentence) - 1, -1, -1):
                            chars = sentence[index]
                            if chars == "" or chars == " ":
                                sentence.pop(index)

                formula_list.sort(key=lambda x: x["coordinates"][1])

                # вставляем формулы в текст
                for non_interact_tokens, formula in zip(page_non_interact_text, formula_list):
                    end_index = find_partial_match_end_index(haystack=original_text, needle=" ".join(non_interact_tokens))
                    original_text = original_text[:end_index] + f"\n$${formula["latex"]}$$\n" + original_text[end_index:]

                try:
                    page_text += original_text.encode('latin1').decode('cp1251')
                except Exception:
                    page_text += original_text

                # распознавание текста на изображениях
                for image_index, img in enumerate(page.get_images(full=True)):
                    try:
                        xref = img[0]
                        base_image = doc.extract_image(xref)
                        image_bytes = base_image["image"]
                        image_stream = BytesIO(image_bytes)
                        image = Image.open(image_stream)
                        recognized_text = reader.image_to_string(image, lang="eng+rus")
                        page_text += recognized_text
                    except Exception as e:
                        logging.warning(f"Error while reading image: {e}")
                        continue

                result_text.append(page_text)

        return ' '.join(result_text)
    except Exception as e:
        raise e


def ultimate_text_reader(documents_bytes: BinaryIO):
    read_doc = documents_bytes.read()
    if read_doc == b'':
        raise HTTPException(status_code=400, detail="Document is empty",
                            headers={"content_type": "text/plain", "tokens": 0})
    try:
        return read_doc.decode('utf-8')
    except UnicodeDecodeError:
        result = chardet.detect(read_doc)
        encoding = result['encoding']
        logging.info(f"Auto detected encoding: {encoding} for plain text file")
        return read_doc.decode(encoding, errors='ignore')


def doc_reader(doc_bytes: BinaryIO) -> str:
    with tempfile.NamedTemporaryFile(delete=True, suffix=".doc") as tmp:
        tmp.write(doc_bytes.read())
        tmp.seek(0)

        document = SpireDocument()
        document.LoadFromFile(tmp.name)

        text = document.GetText()
        return text[71:]


def docx_reader(docx_bytes: BinaryIO) -> str:
    # Загружаем документ из переданного файла
    doc = Document(docx_bytes)

    # Список для хранения текста документа
    doc_text = []

    # Проходим по всем элементам документа
    for element in doc.element.body:
        # Если элемент является параграфом, добавляем его текст в список
        if element.tag.endswith('p'):
            doc_text.append(element.text)
        # Если элемент является таблицей
        elif element.tag.endswith('tbl'):
            # Список для хранения текста таблицы
            table_text = []
            # Проходим по всем строкам таблицы
            for row in element.xpath('./w:tr'):
                # Список для хранения текста текущей строки таблицы
                row_text = []
                # Проходим по всем ячейкам строки
                for cell in row.xpath('./w:tc'):
                    cell_text = cell.xpath('.//w:t')
                    # Собираем текст из всех элементов ячейки
                    cell_text_str = ''.join([t.text for t in cell_text])
                    row_text.append(cell_text_str)
                # Добавляем текст строки в список текста таблицы
                table_text.append(' '.join(row_text))
            # Добавляем весь текст таблицы в основной список текста документа
            doc_text.append('\n'.join(table_text))

    # Возвращаем весь текст документа, преобразованный в строку
    return '\n'.join(doc_text)


def excel_reader(xlsx_bytes: BinaryIO):
    workbook = load_workbook(filename=xlsx_bytes)
    sheet = workbook.active

    data = []
    for row in sheet.iter_rows(values_only=True):
        row_data = [str(cell) for cell in row]
        data.append(' '.join(row_data))

    return "\n".join(data)


def image_reader(image_bytes: BinaryIO):
    image_stream = BytesIO(image_bytes.read())
    image_stream.seek(0)
    image = Image.open(image_stream)

    recognized_text = reader.image_to_string(image, lang="eng+rus")
    return recognized_text


def xls_reader(xls_bytes: BinaryIO):
    workbook = xlrd.open_workbook(file_contents=BytesIO(xls_bytes.read()).getvalue())
    # Выбираем первый лист
    sheet = workbook.sheet_by_index(0)

    data = []
    for row_idx in range(sheet.nrows):
        # Считываем данные по строкам
        row_data = [str(sheet.cell(row_idx, col_idx).value) for col_idx in range(sheet.ncols)]
        data.append(' '.join(row_data))

    return "\n".join(data)


def csv_reader(csv_bytes: BinaryIO):
    csv_read = csv.reader(csv_bytes.read().decode('utf-8').splitlines())

    data = []
    for row in csv_read:
        row_data = [str(cell) for cell in row]
        data.append(' '.join(row_data))

    return "\n".join(data)


def pptx_reader(pptx_bytes: BinaryIO):
    # Загрузите презентацию
    prs = Presentation(BytesIO(pptx_bytes.read()))

    all_text = []  # Список для хранения всего текста

    # Пройдитесь по всем слайдам
    for slide_number, slide in enumerate(prs.slides):
        slide_text = [f"Slide #{slide_number + 1}"]

        # Чтение заголовков и текста из слайдов
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    slide_text.append(run.text)

        # Чтение заметок к слайду, если они есть
        if slide.has_notes_slide:
            notes_slide = slide.notes_slide
            notes_text = notes_slide.notes_text_frame.text
            slide_text.append("Slide notes:")
            slide_text.append(notes_text)

        # Добавляем текст слайда (и заметки, если они есть) в общий список
        all_text.append("\n".join(slide_text))

    # Возвращаем весь текст, собранный из слайдов, объединенный через два переноса строки
    return "\n\n".join(all_text)
